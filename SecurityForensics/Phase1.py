from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

# Create document
doc = Document()

# Title
title = doc.add_heading('LLM Forensics - Phase 1: Pre-training', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Overview
doc.add_heading('Overview', level=1)
doc.add_paragraph(
    'This document covers the pre-training phase of Large Language Models (LLMs), '
    'which forms the foundation for all subsequent capabilities. Understanding this phase '
    'is essential for forensic analysis of LLM behavior, capabilities, and limitations.'
)

# Section 1: Training Data
doc.add_heading('1. Training Data', level=1)

doc.add_heading('1.1 Data Sources', level=2)
data_sources = doc.add_paragraph()
data_sources.add_run('LLMs are trained on massive, diverse text corpora including:\n').bold = True
data_sources.add_run('• Books (fiction and non-fiction)\n')
data_sources.add_run('• Web pages and articles\n')
data_sources.add_run('• Code repositories (GitHub, etc.)\n')
data_sources.add_run('• Wikipedia and encyclopedic content\n')
data_sources.add_run('• Scientific papers\n')
data_sources.add_run('• Social media content\n')
data_sources.add_run('• News articles\n')
data_sources.add_run('• Historical archives')

doc.add_heading('1.2 Data Characteristics', level=2)
data_char = doc.add_paragraph()
data_char.add_run('Key characteristics of training data:\n').bold = True
data_char.add_run('• Scale: trillions of tokens\n')
data_char.add_run('• Diversity: multiple languages, domains, formats\n')
data_char.add_run('• Quality variance: includes both high-quality and noisy content\n')
data_char.add_run('• Biases: reflects human biases in source material\n')
data_char.add_run('• Temporal coverage: limited to data available before training cutoff')

# Section 2: Training Objective
doc.add_heading('2. Training Objective', level=1)

doc.add_paragraph(
    'The primary objective during pre-training is to minimize the loss on next-token prediction. '
    'This is known as autoregressive language modeling.'
)

doc.add_heading('2.1 Mathematical Foundation', level=2)
math_para = doc.add_paragraph()
math_para.add_run('The model learns to maximize:\n').bold = True
math_para.add_run('L = Σ log P(w_t | w_1, w_2, ..., w_{t-1})')
math_para.add_run('\n\nWhere:\n').bold = True
math_para.add_run('• w_t = the t-th token\n')
math_para.add_run('• w_1 to w_{t-1} = previous tokens (context)\n')
math_para.add_run('• P = probability distribution over vocabulary')

doc.add_heading('2.2 Tokenization', level=2)
token_doc = doc.add_paragraph()
token_doc.add_run('Before training, text is converted to tokens using:\n').bold = True
token_doc.add_run('• Byte Pair Encoding (BPE)\n')
token_doc.add_run('• WordPiece\n')
token_doc.add_run('• SentencePiece\n')
token_doc.add_run('\nThis creates a vocabulary of ~30,000-50,000 tokens.')

# Section 3: Architecture
doc.add_heading('3. Model Architecture', level=1)

doc.add_heading('3.1 Transformer Base', level=2)
arch_para = doc.add_paragraph()
arch_para.add_run('Modern LLMs use the Transformer architecture introduced in "Attention Is All You Need" (2017).\n').bold = True
arch_para.add_run('Key components:\n')
arch_para.add_run('• Multi-head self-attention\n')
arch_para.add_run('• Feed-forward neural networks\n')
arch_para.add_run('• Layer normalization\n')
arch_para.add_run('• Positional encodings\n')
arch_para.add_run('• Residual connections')

doc.add_heading('3.2 Scale', level=2)
scale_para = doc.add_paragraph()
scale_para.add_run('LLM scale includes:\n').bold = True
scale_para.add_run('• Parameters: billions to hundreds of billions\n')
scale_para.add_run('• Layers: dozens to hundreds of transformer blocks\n')
scale_para.add_run('• Attention heads: dozens per layer\n')
scale_para.add_run('• Context length: 1K to 100K+ tokens')

# Section 4: Training Process
doc.add_heading('4. Training Process', level=1)

doc.add_heading('4.1 Optimization', level=2)
opt_para = doc.add_paragraph()
opt_para.add_run('Training uses:\n').bold = True
opt_para.add_run('• Optimizer: AdamW (most common)\n')
opt_para.add_run('• Learning rate scheduling: warmup + decay\n')
opt_para.add_run('• Batch size: large (thousands to millions of tokens)\n')
opt_para.add_run('• Mixed precision training (FP16/BF16)\n')
opt_para.add_run('• Gradient clipping to prevent instability')

doc.add_heading('4.2 Hardware Requirements', level=2)
hw_para = doc.add_paragraph()
hw_para.add_run('Pre-training requires massive compute:\n').bold = True
hw_para.add_run('• Thousands of GPUs (A100, H100, etc.)\n')
hw_para.add_run('• Training time: weeks to months\n')
hw_para.add_run('• Estimated cost: millions of dollars for largest models')

# Section 5: What the Model Learns
doc.add_heading('5. What the Model Learns', level=1)

doc.add_heading('5.1 Knowledge Compression', level=2)
doc.add_paragraph(
    'Pre-training compresses information from training data into model weights. '
    'The model does not store facts explicitly but learns patterns, relationships, '
    'and statistical regularities in language.'
)

doc.add_heading('5.2 Capabilities Acquired', level=2)
cap_para = doc.add_paragraph()
cap_para.add_run('After pre-training, models acquire:\n').bold = True
cap_para.add_run('• Language understanding and generation\n')
cap_para.add_run('• Grammar and syntax\n')
cap_para.add_run('• Reasoning patterns\n')
cap_para.add_run('• World knowledge (limited by training data)\n')
cap_para.add_run('• Coding abilities\n')
cap_para.add_run('• Cultural and social biases\n')
cap_para.add_run('• Stylistic patterns from training data')

doc.add_heading('5.3 Limitations Introduced', level=2)
lim_para = doc.add_paragraph()
lim_para.add_run('Pre-training introduces fundamental limitations:\n').bold = True
lim_para.add_run('• Hallucinations: confident but incorrect outputs\n')
lim_para.add_run('• Knowledge cutoff: no awareness of events after training\n')
lim_para.add_run('• Biases: reflects training data biases\n')
lim_para.add_run('• Factual errors: not a reliable fact database\n')
lim_para.add_run('• No verification: cannot check external truth')

# Section 6: Forensic Implications
doc.add_heading('6. Forensic Implications', level=1)

doc.add_heading('6.1 Evidence of Pre-training', level=2)
forensic_para = doc.add_paragraph()
forensic_para.add_run('Forensic analysts can identify pre-training effects through:\n').bold = True
forensic_para.add_run('• Memorization of training data (especially for short texts)\n')
forensic_para.add_run('• Stylistic patterns from specific domains\n')
forensic_para.add_run('• Knowledge boundaries matching training cutoff\n')
forensic_para.add_run('• Bias patterns reflecting source data demographics\n')
forensic_para.add_run('• Language capabilities matching training corpora')

doc.add_heading('6.2 Analysis Methods', level=2)
methods_para = doc.add_paragraph()
methods_para.add_run('Black-box forensic methods:\n').bold = True
methods_para.add_run('• Probing for specific known facts from training period\n')
methods_para.add_run('• Testing knowledge cutoff dates\n')
methods_para.add_run('• Measuring bias through structured prompts\n')
methods_para.add_run('• Identifying memorization through exact match queries\n')
methods_para.add_run('• Benchmarking on domain-specific tasks')

doc.add_heading('6.3 Red-teaming Findings', level=2)
doc.add_paragraph(
    'Red-teaming has revealed that pre-training captures harmful content from the internet, '
    'including misinformation, hate speech, and dangerous instructions. This is why '
    'subsequent fine-tuning phases focus on alignment and safety.'
)

# Section 7: Summary
doc.add_heading('7. Summary', level=1)
summary = doc.add_paragraph()
summary.add_run('Key Takeaways:\n').bold = True
summary.add_run('1. Pre-training forms the foundation of all LLM capabilities\n')
summary.add_run('2. Knowledge is compressed into patterns in model weights\n')
summary.add_run('3. Training data determines what the model can know and express\n')
summary.add_run('4. Pre-training introduces fundamental limitations (hallucinations, biases)\n')
summary.add_run('5. Forensic analysis can probe these characteristics systematically\n')
summary.add_run('6. Fine-tuning (Phase 2) modifies but does not erase pre-training effects')

# Add a footer note
doc.add_paragraph('\n' + '='*50)
doc.add_paragraph('Document created for LLM forensics research purposes.')
doc.add_paragraph('Based on Phase 1: Pre-training information.')

# Save document
doc.save('LLM_Forensics_Phase1_Pretraining.docx')
print("Document created: LLM_Forensics_Phase1_Pretraining.docx")
